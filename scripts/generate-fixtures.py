import math,struct,wave,subprocess
from pathlib import Path
from PIL import Image,ImageDraw
from docx import Document
from pptx import Presentation
from openpyxl import Workbook
from odf.opendocument import OpenDocumentSpreadsheet
from odf.table import Table,TableRow,TableCell
from odf.text import P
root=Path(__file__).resolve().parents[1]
out=root/'tests/fixtures/corpus';out.mkdir(parents=True,exist_ok=True)
(out/'עברית.txt').write_text('ארכיטקטורת SparkMoE מאפשרת חיפוש מקומי במחשב. הפגישה ביום רביעי בוטלה.\n',encoding='utf-8')
(out/'ocean.txt').write_text('A child is playing with a ball on the beach near the ocean.',encoding='utf-8')
(out/'notes.md').write_text('# Search architecture\n\nSQLite stores metadata and full text. USearch stores embeddings.\n',encoding='utf-8')
(out/'auth.rs').write_text('/// Connect to OAuth locally\nfn connect_oauth(token: &str) -> bool { !token.is_empty() }\n',encoding='utf-8')
(out/'opaque.bin').write_bytes(bytes(range(256)))
(out/'.env').write_text('SYNTHETIC_SECRET=must_be_excluded',encoding='utf-8')
doc=Document();doc.add_heading('Meeting cancellation',0);doc.add_paragraph('The Wednesday meeting was cancelled.');doc.save(out/'meeting.docx')
deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[1]);slide.shapes.title.text='Local search architecture';slide.placeholders[1].text='SQLite metadata with hybrid semantic search';deck.save(out/'architecture.pptx')
workbook=Workbook();sheet=workbook.active;sheet.title='Performance';sheet.append(['Model','Latency']);sheet.append(['Synthetic',12]);workbook.save(out/'performance.xlsx')
ods=OpenDocumentSpreadsheet();table=Table(name='Project');row=TableRow();cell=TableCell(valuetype='string');cell.addElement(P(text='Offline local search'));row.addElement(cell);table.addElement(row);ods.spreadsheet.addElement(table);ods.save(str(out/'notes.ods'))
image=Image.new('RGB',(640,360),(170,218,247));d=ImageDraw.Draw(image);d.rectangle((0,160,640,260),fill=(37,121,195));d.rectangle((0,260,640,360),fill=(235,210,150));d.ellipse((480,40,530,90),fill=(249,224,74));d.ellipse((250,235,280,265),fill=(210,70,55));image.save(out/'beach.png')
with wave.open(str(out/'tone.wav'),'wb') as audio:
    audio.setnchannels(1);audio.setsampwidth(2);audio.setframerate(16000);audio.writeframes(b''.join(struct.pack('<h',int(5000*math.sin(i*2*math.pi*440/16000))) for i in range(16000*3)))
objects=[b'<< /Type /Catalog /Pages 2 0 R >>',b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>']
stream=b'BT /F1 18 Tf 60 720 Td (Offline search with SQLite and embeddings) Tj ET'
objects.append(b'<< /Length '+str(len(stream)).encode()+b' >>\nstream\n'+stream+b'\nendstream')
data=bytearray(b'%PDF-1.4\n');offsets=[0]
for i,obj in enumerate(objects,1):offsets.append(len(data));data.extend(f'{i} 0 obj\n'.encode()+obj+b'\nendobj\n')
xref=len(data);data.extend(f'xref\n0 {len(objects)+1}\n0000000000 65535 f \n'.encode());data.extend(b''.join(f'{o:010} 00000 n \n'.encode() for o in offsets[1:]));data.extend(f'trailer << /Size {len(objects)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF'.encode());(out/'search.pdf').write_bytes(data)
ffmpeg=root/'resources/ffmpeg/bin/ffmpeg.exe'
if ffmpeg.exists():subprocess.run([str(ffmpeg),'-v','error','-y','-loop','1','-i',str(out/'beach.png'),'-i',str(out/'tone.wav'),'-t','3','-c:v','libopenh264','-pix_fmt','yuv420p','-c:a','aac',str(out/'beach.mp4')],check=True,creationflags=subprocess.CREATE_NO_WINDOW)
print('Synthetic fixtures generated')
